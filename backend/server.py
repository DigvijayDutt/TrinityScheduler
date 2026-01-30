from fastapi import FastAPI, HTTPException, BackgroundTasks
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import StreamingResponse
from fastapi import Depends
import random
from psycopg2.pool import SimpleConnectionPool
from jose import jwt
from datetime import datetime, timedelta, date
import pandas as pd
import io
import os
from psycopg2.pool import SimpleConnectionPool
from psycopg2 import sql
from psycopg2 import errors
from fastapi.responses import Response
from pydantic import BaseModel
from typing import Optional
from openpyxl.styles import numbers
from config import settings
from passlib.context import CryptContext


# 🔔 EMAIL SERVICE
from email_service import send_assignment_email

# --------------------------
# Config
# --------------------------

# prabhat's change
class SkillUpdate(BaseModel):
    name: str
    description: Optional[str] = ""



SECRET_KEY = "SECRET_KEY"
ALGORITHM = "HS256"

app = FastAPI()

# --------------------------
# CORS
# --------------------------
app.add_middleware(
    CORSMiddleware,
    allow_origins=settings.CORS_ORIGINS.split(","),
    allow_methods=["*"],
    allow_headers=["*"],
)

# --------------------------
# PostgreSQL Pool
# --------------------------
pool = SimpleConnectionPool(
    minconn=1,
    maxconn=10,
    host=settings.DATABASE_HOST,
    database=settings.DATABASE_NAME,
    user=settings.DATABASE_USER,
    password=settings.DATABASE_PASSWORD,
    port=settings.DATABASE_PORT
)

def get_db():
    conn = pool.getconn()
    try:
        yield conn
    finally:
        pool.putconn(conn)

pwd_context = CryptContext(schemes=["bcrypt"], deprecated="auto")

def verify_password(plain, hashed):
    return pwd_context.verify(plain, hashed)

@app.get("/health")
def health():
    return {"status": "ok"}


# --------------------------
# LOGIN
# --------------------------
@app.post("/login")
def login(data: dict):
    conn = pool.getconn()
    try:
        cur = conn.cursor()
        email = data.get("email")
        password = data.get("password")

        cur.execute(
            "SELECT id, username, email, password_hash FROM users WHERE email = %s",
            (email,)
        )
        row = cur.fetchone()

        if not row:
            raise HTTPException(status_code=401, detail="User not found")

        user_id, db_username, email, password_hash = row

        if password != password_hash:
            raise HTTPException(status_code=401, detail="Invalid password")

        token = jwt.encode(
            {
                "id": user_id,
                "username": db_username,
                "exp": datetime.utcnow() + timedelta(hours=1)
            },
            SECRET_KEY,
            algorithm=ALGORITHM
        )

        return {"message": "Login successful", "token": token}

    finally:
        cur.close()
        pool.putconn(conn)

# --------------------------
# JOB TYPES
# --------------------------
@app.get("/jobTypes")
def getJobType():
    conn = pool.getconn()
    try:
        cur = conn.cursor()
        cur.execute("SELECT type FROM jobs")
        rows = cur.fetchall()
        return [row[0] for row in rows]
    finally:
        cur.close()
        pool.putconn(conn)

@app.get("/employeeNames")
def getEmployeeName():
    conn = pool.getconn()
    try:
        cur = conn.cursor()
        cur.execute("SELECT name FROM employees")
        return [row[0] for row in cur.fetchall()]
    finally:
        cur.close()
        pool.putconn(conn)

@app.get("/employees")
def getEmployees():
    conn = pool.getconn()
    try:
        cur = conn.cursor()
        cur.execute("SELECT * FROM employees ORDER BY id ASC")
        rows = cur.fetchall()
        colnames = [desc[0] for desc in cur.description]
        return [dict(zip(colnames, row)) for row in rows]
    finally:
        cur.close()
        pool.putconn(conn)



# Digvijay's change
@app.post("/employees")
def createEmployee(emp: dict):
    conn = pool.getconn()
    try:
        cur = conn.cursor()

        if "name" not in emp or "email" not in emp:
            raise HTTPException(status_code=400, detail="name and email are required")

        columns = []
        values = []

        for key, value in emp.items():
            columns.append(key)
            values.append(value)

        column_names = ", ".join(columns)
        placeholders = ", ".join(["%s"] * len(values))

        query = f"""
            INSERT INTO employees ({column_names})
            VALUES ({placeholders})
            RETURNING *
        """

        cur.execute(query, tuple(values))
        row = cur.fetchone()
        conn.commit()

        cols = [desc[0] for desc in cur.description]
        return dict(zip(cols, row))

    finally:
        cur.close()
        pool.putconn(conn)




# prabhat's delete
@app.delete("/employees/{emp_id}")
def deleteEmployee(emp_id: int):
    conn = pool.getconn()
    try:
        cur = conn.cursor()
        cur.execute("DELETE FROM employees WHERE id = %s RETURNING *", (emp_id,))
        row = cur.fetchone()
        if not row:
            raise HTTPException(status_code=404, detail="Employee not found")
        conn.commit()
        return {"message": "Employee deleted successfully", "id": emp_id}
    finally:
        cur.close()
        pool.putconn(conn)


# @app.get("/skills")
# def getSkills():
#     conn = pool.getconn()
#     try:
#         cur = conn.cursor()
#         cur.execute("SELECT * FROM skills")
#         return [row[0] for row in cur.fetchall()]
#     finally:
#         cur.close()
#         pool.putconn(conn)


# prabhat's
@app.get("/skills")
def getSkills():
    conn = pool.getconn()
    try:
        cur = conn.cursor()
        cur.execute("SELECT id, onsite_work, description FROM skills ORDER BY id")
        rows = cur.fetchall()
        return [{"id": r[0], "name": r[1], "description": r[2]} for r in rows]
    finally:
        cur.close()
        pool.putconn(conn)


# --------------------------
# JOBS
# --------------------------
@app.get("/jobs")
def getJobs():
    conn = pool.getconn()
    try:
        cur = conn.cursor()
        cur.execute("SELECT * FROM jobs")
        rows = cur.fetchall()
        colnames = [desc[0] for desc in cur.description]
        return [dict(zip(colnames, row)) for row in rows]
    finally:
        cur.close()
        pool.putconn(conn)

@app.get("/jobs/normalized")
def getJobsNormalized():
    conn = pool.getconn()
    try:
        cur = conn.cursor()

        # fetch all skill column names
        cur.execute("SELECT onsite_work FROM skills ORDER BY id")
        skill_names = [r[0] for r in cur.fetchall()]

        # fetch all jobs
        cur.execute("SELECT * FROM jobs")
        rows = cur.fetchall()
        colnames = [desc[0] for desc in cur.description]

        jobs = []
        for row in rows:
            job = dict(zip(colnames, row))

            # extract skills dynamically
            skills = {
                k: job[k]
                for k in skill_names
                if job.get(k, 0) and job[k] > 0
            }

            # remove skill columns from root
            for k in skill_names:
                job.pop(k, None)

            job["skills"] = skills
            jobs.append(job)

        return jobs
    finally:
        cur.close()
        pool.putconn(conn)



@app.get("/scheduledjobs")
def getScheduledJobs():
    conn = pool.getconn()
    try:
        cur = conn.cursor()
        # cur.execute("SELECT * FROM scheduledjobs")
        cur.execute("SELECT * FROM scheduledjobs ORDER BY start_date")
        # rows = cur.fetchall()
        # colnames = [desc[0] for desc in cur.description]
        # return [dict(zip(colnames, row)) for row in rows]

        rows = cur.fetchall()
        colnames = [desc[0] for desc in cur.description]

        jobs = []

        for row in rows:
            job = dict(zip(colnames, row))
            assigned_ids = job.get("assigned", [])

            if not assigned_ids:
                job["assigned_staff_display"] = ""
                jobs.append(job)
                continue

            # fetch assigned employees
            cur.execute("""
                SELECT id, name, teamlead
                FROM employees
                WHERE id = ANY(%s)
            """, (assigned_ids,))

            emps = cur.fetchall()

            # find highest teamlead
            max_tl = max(e[2] for e in emps)
            teamlead_id = next(e[0] for e in emps if e[2] == max_tl)

            names = []
            for emp_id, name, tl in emps:
                if emp_id == teamlead_id:
                    names.append(f"{name}(Team lead)")
                else:
                    names.append(name)

            job["assigned_staff_display"] = ", ".join(names)
            jobs.append(job)

        return jobs



    finally:
        cur.close()
        pool.putconn(conn)

# --------------------------
# DOWNLOAD SCHEDULED JOBS
# --------------------------
@app.get("/scheduledjobs/export")
def download_scheduled_jobs():
    conn = pool.getconn()
    try:
        query = """
            SELECT
                sj.id,
                sj.address,
                sj.type,
                sj.client,
                sj.status,
                sj.start_date,
                COALESCE(
                    STRING_AGG(
                        CASE
                            WHEN e.id = tl.id
                            THEN e.name || ' (Team lead)'
                            ELSE e.name
                        END,

                        E'\n'
                        ORDER BY e.teamlead DESC, e.name
                    ),
                    ''
                ) AS assigned
            FROM scheduledjobs sj
            LEFT JOIN LATERAL unnest(sj.assigned) AS emp_id ON TRUE
            LEFT JOIN employees e ON e.id = emp_id
            LEFT JOIN LATERAL (
                SELECT id
                FROM employees
                WHERE id = ANY(sj.assigned)
                ORDER BY teamlead DESC, id ASC
                LIMIT 1
            ) tl ON TRUE


            GROUP BY
                sj.id,
                sj.address,
                sj.type,
                sj.client,
                sj.status,
                sj.start_date
            ORDER BY sj.id;
        """

        df = pd.read_sql(query, conn)
        df["start_date"] = pd.to_datetime(df["start_date"])

        output = io.BytesIO()

        from openpyxl.styles import Font, PatternFill, Alignment
        from openpyxl.utils import get_column_letter

        with pd.ExcelWriter(output, engine="openpyxl") as writer:
            df.to_excel(writer, index=False, sheet_name="Scheduled Jobs")
            worksheet = writer.sheets["Scheduled Jobs"]

            # -------------------------
            # HEADER STYLE
            # -------------------------
            header_fill = PatternFill("solid", fgColor="1F4E78")
            header_font = Font(color="FFFFFF", bold=True)

            for cell in worksheet[1]:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center", vertical="center")

            # -------------------------
            # FILTERS + FREEZE HEADER
            # -------------------------
            worksheet.auto_filter.ref = worksheet.dimensions
            worksheet.freeze_panes = "A2"

            # -------------------------
            # AUTO COLUMN WIDTH
            # -------------------------
            for col_idx, col in enumerate(df.columns, start=1):
                max_length = max(
                    df[col].astype(str).map(len).max(),
                    len(col)
                )
                worksheet.column_dimensions[
                    get_column_letter(col_idx)
                ].width = min(max_length + 3, 40)

            # -------------------------
            # DATE FORMAT
            # -------------------------
            start_date_col = df.columns.get_loc("start_date") + 1
            for row in range(2, len(df) + 2):
                worksheet.cell(row=row, column=start_date_col).number_format = "DD-MM-YYYY"

            # -------------------------
            # ASSIGNED STAFF SPACING
            # -------------------------
            assigned_col = df.columns.get_loc("assigned") + 1
            for row in range(2, len(df) + 2):
                worksheet.row_dimensions[row].height = 35
                worksheet.cell(row=row, column=assigned_col).alignment = Alignment(
                    wrap_text=True,
                    vertical="top"
                )

        output.seek(0)

        return Response(
            content=output.getvalue(),
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={
                "Content-Disposition": "attachment; filename=scheduled_jobs.xlsx"
            }
        )

    finally:
        pool.putconn(conn)



# @app.get("/scheduledjobs/{job_id}")
# def getScheduledJobById(job_id: str):
#     conn = pool.getconn()
#     try:
#         cur = conn.cursor()
#         cur.execute(
#             "SELECT * FROM scheduledjobs WHERE id = %s",
#             (job_id,)
#         )
#         row = cur.fetchone()

#         if not row:
#             raise HTTPException(status_code=404, detail="Job not found")

#         colnames = [desc[0] for desc in cur.description]
#         return dict(zip(colnames, row))

#     finally:
#         cur.close()
#         pool.putconn(conn)

@app.get("/scheduledjobs/{job_id}")
def getScheduledJobById(job_id: str):
    conn = pool.getconn()
    try:
        cur = conn.cursor()

        cur.execute(
            "SELECT * FROM scheduledjobs WHERE id = %s",
            (job_id,)
        )
        row = cur.fetchone()

        if not row:
            raise HTTPException(status_code=404, detail="Job not found")

        colnames = [desc[0] for desc in cur.description]
        job = dict(zip(colnames, row))

        assigned_ids = job.get("assigned", [])

        if not assigned_ids:
            job["assigned_staff_display"] = ""
            job["team_lead_id"] = None
            return job

        # fetch assigned employees
        cur.execute("""
            SELECT id, name, teamlead
            FROM employees
            WHERE id = ANY(%s)
        """, (assigned_ids,))

        emps = cur.fetchall()

        # find highest teamlead (SAME AS JOB PAGE)
        max_tl = max(e[2] for e in emps)
        teamlead_id = next(e[0] for e in emps if e[2] == max_tl)

        names = []
        for emp_id, name, tl in emps:
            if emp_id == teamlead_id:
                names.append(f"{name}(Team lead)")
            else:
                names.append(name)

        job["assigned_staff_display"] = ", ".join(names)
        job["team_lead_id"] = teamlead_id

        return job

    finally:
        cur.close()
        pool.putconn(conn)



# @app.post("/scheduledjobs")
# def postScheduledJobs(data: dict, background_tasks: BackgroundTasks):
#     conn = pool.getconn()
#     try:
#         cur = conn.cursor()

#         job_id = f"J-{random.randint(10000,99999)}"
#         address = data.get("address")
#         job_type = data.get("jobType")
#         client = data.get("client")
#         assigned = [int(e) for e in (data.get("assigned") or [])]
#         start_date = datetime.fromisoformat(data.get("time"))

#         cur.execute("""
#             INSERT INTO scheduledjobs
#             (id, address, type, client, status, assigned, start_date)
#             VALUES (%s, %s, %s, %s, %s, %s::int[], %s)
#         """, (job_id, address, job_type, client, "Scheduled", assigned, start_date))

#         employee_names = []
#         employee_emails = []

#         for emp_id in assigned:
#             cur.execute(
#                 "INSERT INTO assignments (jobid, empid, jobdate) VALUES (%s, %s, %s)",
#                 (job_id, emp_id, start_date)
#             )

#             cur.execute(
#                 "SELECT name, email FROM employees WHERE id = %s",
#                 (emp_id,)
#             )
#             emp = cur.fetchone()
#             if emp and emp[1]:
#                 employee_names.append(emp[0])
#                 employee_emails.append(emp[1])

#         if employee_emails:
#             background_tasks.add_task(
#                 send_assignment_email,
#                 employee_emails,
#                 employee_names,
#                 job_id,
#                 job_type,
#                 start_date,
#                 client,
#                 address
#             )

#         conn.commit()
#         return {"message": "Job created", "id": job_id}

#     finally:
#         cur.close()
#         pool.putconn(conn)


# prabhat's chnage
@app.post("/scheduledjobs")
def postScheduledJobs(data: dict, background_tasks: BackgroundTasks):
    conn = pool.getconn()
    try:
        cur = conn.cursor()

        job_id = f"J-{random.randint(10000,99999)}"
        address = data.get("address")
        job_type = data.get("jobType")
        client = data.get("client")
        assigned = [int(e) for e in (data.get("assigned") or [])]
        start_date = datetime.fromisoformat(data.get("time"))
        
        loss_type = data.get("lossType")
        project_manager = data.get("projectManager")
        vehicle = data.get("vehicle")
        special_instructions = data.get("specialInstructions")

        job_time = None
        if data.get("time"):
            job_time = datetime.fromisoformat(data["time"]).date()

        # Insert main job record
        cur.execute("""
            INSERT INTO scheduledjobs (
                id,
                address,
                type,
                client,
                status,
                assigned,
                start_date,
                loss_type,
                project_manager,
                vehicle,
                special_instructions,
                job_time
            )
            VALUES (%s, %s, %s, %s, %s, %s::int[], %s, %s, %s, %s, %s, %s)
        """, (
            job_id,
            address,
            job_type,
            client,
            "Scheduled",
            assigned,
            start_date,
            loss_type,
            project_manager,
            vehicle,
            special_instructions,
            job_time
        ))


        employee_names = []
        employee_emails = []
        team_lead_name = None
        teamlead_level = -1


        # Assign employees safely
        # for emp_id in assigned:
        #     try:
        #         cur.execute(
        #             "INSERT INTO assignments (jobid, empid, jobdate) VALUES (%s, %s, %s)",
        #             (job_id, emp_id, start_date)
        #         )
        #     except errors.UniqueViolation:
        #         conn.rollback()
        #         # fetch employee name
        #         cur.execute("SELECT name FROM employees WHERE id = %s", (emp_id,))
        #         emp_name = cur.fetchone()[0]
        #         raise HTTPException(
        #             status_code=400,
        #             detail=f"Employee {emp_name} is already assigned on {start_date.date()}"
        #         )




        for emp_id in assigned:
            # Check if employee has a conflicting job on this date
            cur.execute("""
                SELECT a.jobid, sj.status
                FROM assignments a
                JOIN scheduledjobs sj ON sj.id = a.jobid
                WHERE a.empid = %s AND a.jobdate::date = %s
            """, (emp_id, start_date.date()))
            
            existing = cur.fetchone()

            if existing:
                jobid, status = existing
                if status.lower().strip() in ('scheduled', 'in progress'):
                    # Conflict! Cannot assign
                    cur.execute("SELECT name FROM employees WHERE id = %s", (emp_id,))
                    emp_name = cur.fetchone()[0]
                    raise HTTPException(
                        status_code=400,
                        detail=f"Employee {emp_name} is already assigned on {start_date.date()}"
                    )
                else:
                    # Already has assignment but Completed → update the assignment to new job
                    cur.execute(
                        "UPDATE assignments SET jobid=%s WHERE empid=%s AND jobdate=%s",
                        (job_id, emp_id, start_date)
                    )
            else:
                # No assignment yet → insert normally
                cur.execute(
                    "INSERT INTO assignments (jobid, empid, jobdate) VALUES (%s, %s, %s)",
                    (job_id, emp_id, start_date)
                )



        
        

        # conflicts = []

        # for emp_id in assigned:
        #     try:
        #         cur.execute(
        #             "INSERT INTO assignments (jobid, empid, jobdate) VALUES (%s, %s, %s)",
        #             (job_id, emp_id, start_date)
        #         )
        #     except errors.UniqueViolation:
        #         conn.rollback()
        #         cur.execute("SELECT name FROM employees WHERE id = %s", (emp_id,))
        #         emp_name = cur.fetchone()[0]
        #         conflicts.append(emp_name)

        # if conflicts:
        #     names_str = ", ".join(conflicts)
        #     raise HTTPException(
        #         status_code=400,
        #         detail=f"Employees {names_str} are already assigned on {start_date.date()}"
        #     )


            # fetch email for sending assignment
            cur.execute(
                "SELECT name, email, teamlead FROM employees WHERE id = %s",
                (emp_id,)
            )
            emp = cur.fetchone()

            if emp:
                name, email, teamlead = emp

                employee_names.append(name)
                if email:
                    employee_emails.append(email)

                # 👇 SAME LOGIC AS JOB PAGE
                if teamlead is not None and teamlead > teamlead_level:
                    teamlead_level = teamlead
                    team_lead_name = name


        if employee_emails:
            background_tasks.add_task(
                send_assignment_email,
                employee_emails,
                employee_names,
                team_lead_name,   # 👈 ADD THIS
                job_id,
                job_type,
                start_date,
                client,
                address,
                loss_type,
                project_manager,
                vehicle,
                special_instructions
            )


        conn.commit()
        return {"message": "Job created", "id": job_id}

    finally:
        cur.close()
        pool.putconn(conn)



@app.delete("/scheduledjobs/{job_id}")
def deleteJobs(job_id: str):
    conn = pool.getconn()
    try:
        cur = conn.cursor()
        cur.execute("DELETE FROM scheduledjobs WHERE id = %s", (job_id,))
        conn.commit()
        return {"message": "Job deleted", "id": job_id}
    finally:
        cur.close()
        pool.putconn(conn)

@app.put("/scheduledjobs/{job_id}")
def editJobs(job_id: str, data: dict):
    conn = pool.getconn()
    try:
        cur = conn.cursor()
        cur.execute(
            "SELECT address, type, client, status, start_date FROM scheduledjobs WHERE id = %s",
            (job_id,)
        )
        row = cur.fetchone()
        if not row:
            raise HTTPException(status_code=404, detail="Job not found")

        address, job_type, client, status, start_date = row

        cur.execute(
            """
            UPDATE scheduledjobs
            SET address = %s, type = %s, client = %s, status = %s, start_date = %s
            WHERE id = %s
            """,
            (
                data.get("address", address),
                data.get("type", job_type),
                data.get("client", client),
                data.get("status", status),
                data.get("start_date", start_date),
                job_id
            )
        )

        conn.commit()
        return {"message": "Job updated", "id": job_id}

    finally:
        cur.close()
        pool.putconn(conn)

# --------------------------
# CALENDAR
# --------------------------
@app.get("/calendar/{Date}")
def getCalendar(Date: date):
    conn = pool.getconn()
    try:
        cur = conn.cursor()
        cur.execute("""
            SELECT sj.start_date, sj.id AS job_id, sj.type,
                   e.name AS employee
            FROM scheduledjobs sj
            JOIN assignments ja ON ja.jobid = sj.id
            JOIN employees e ON e.id = ja.empid
            WHERE sj.start_date::date = %s
            ORDER BY sj.start_date;
        """, (Date,))
        rows = cur.fetchall()
        cols = [desc[0] for desc in cur.description]
        return [dict(zip(cols, row)) for row in rows]
    finally:
        cur.close()
        pool.putconn(conn)

# @app.get("/calendar")
# def getAllCalendar():
#     conn = pool.getconn()
#     try:
#         cur = conn.cursor()
#         cur.execute("SELECT * FROM assignments")
#         rows = cur.fetchall()
#         cols = [desc[0] for desc in cur.description]
#         return [dict(zip(cols, row)) for row in rows]
#     finally:
#         cur.close()
#         pool.putconn(conn)

# prabhat's
@app.get("/calendar")
def getAllCalendar():
    conn = pool.getconn()
    try:
        cur = conn.cursor()
        cur.execute("""
            SELECT
                sj.id AS jobid,
                sj.start_date::date AS jobdate,
                sj.status
            FROM scheduledjobs sj
            ORDER BY sj.start_date;
        """)
        rows = cur.fetchall()
        cols = [desc[0] for desc in cur.description]
        return [dict(zip(cols, row)) for row in rows]
    finally:
        cur.close()
        pool.putconn(conn)


# @app.put('/skills/{skill},{old}')
# def editSkill(skill: str, old: str):
#     conn = pool.getconn()
#     try:
#         cur = conn.cursor()
#         cur.execute("update skills set onsite_work = %s where onsite_work = %s",(skill,old))
#     finally:
#         cur.close()
#         pool.putconn(conn)


# prabhat's
@app.put("/skills/{skill_id}")
def updateSkill(skill_id: int, data: SkillUpdate):
    conn = pool.getconn()
    try:
        cur = conn.cursor()
        try:
            cur.execute(
                "UPDATE skills SET onsite_work = %s, description = %s WHERE id = %s",
                (data.name, data.description, skill_id)
            )
            if cur.rowcount == 0:
                raise HTTPException(status_code=404, detail="Skill not found")

            conn.commit()
            return {"message": "Skill updated", "id": skill_id, "name": data.name}

        except errors.UniqueViolation:
            conn.rollback()
            raise HTTPException(
                status_code=400,
                detail="Skill already exists"
            )
    finally:
        cur.close()
        pool.putconn(conn)





# @app.post('/skills')
# def addSkill(data: dict):
#     conn = pool.getconn()
#     skillname = data.get("skillName")
#     try:
#         cur = conn.cursor()
#         cur.execute("insert into skills (onsite_work) values (%s)", (skillname,))
#         conn.commit()
#         return {"message": "skill created", "name": skillname}
#     finally:
#         cur.close()
#         pool.putconn(conn)


# prabhat's
@app.post("/skills")
def addSkill(data: dict):
    conn = pool.getconn()
    skillname = data.get("skillName")
    description = data.get("description", "")
    try:
        cur = conn.cursor()
        try:
            cur.execute(
                "INSERT INTO skills (onsite_work, description) VALUES (%s, %s) RETURNING id",
                (skillname, description)
            )
            skill_id = cur.fetchone()[0]
            query = sql.SQL(
                "ALTER TABLE jobs ADD COLUMN {} INTEGER"
            ).format(sql.Identifier(skillname))
            cur.execute(query)
            query1 = sql.SQL(
                "ALTER TABLE employees ADD COLUMN {} INTEGER"
            ).format(sql.Identifier(skillname))
            cur.execute(query1)
            conn.commit()
            return {"id": skill_id, "name": skillname, "description": description}

        except errors.UniqueViolation:
            conn.rollback()
            raise HTTPException(
                status_code=400,
                detail="Skill already exists"
            )
    finally:
        cur.close()
        pool.putconn(conn)





# @app.delete("/skills/{id}")
# def deleteJobs(id: str):
#     conn = pool.getconn()
#     try:
#         cur = conn.cursor()
#         cur.execute("DELETE FROM skills WHERE onsite_work = %s", (id,))
#         conn.commit()
#         return {"message": "skill deleted", "id": id}
#     finally:
#         cur.close()
#         pool.putconn(conn)


# prabhat's
@app.delete("/skills/{skill_id}")
def deleteSkill(skill_id: int):
    conn = pool.getconn()
    try:
        cur = conn.cursor()
        cur.execute(
            "DELETE FROM skills WHERE id = %s RETURNING id",
            (skill_id,)
        )
        row = cur.fetchone()
        if not row:
            raise HTTPException(status_code=404, detail="Skill not found")

        conn.commit()
        return {"message": "Skill deleted", "id": skill_id}
    finally:
        cur.close()
        pool.putconn(conn)


# @app.get("/busystaff")
# def getBusyStaff():
#     conn = pool.getconn()
#     try:
#         cur = conn.cursor()
#         cur.execute("select empid from assignments")
#         rows = cur.fetchall()
#         return [row[0] for row in rows]
#     finally:
#         cur.close()
#         pool.putconn(conn)

@app.get("/busystaff")
def getBusyStaff():
    conn = pool.getconn()
    try:
        cur = conn.cursor()
        cur.execute("""
            SELECT DISTINCT a.empid
            FROM assignments a
            JOIN scheduledjobs sj ON sj.id = a.jobid
            WHERE a.jobdate::date = CURRENT_DATE
            AND LOWER(TRIM(sj.status)) IN ('scheduled', 'in progress')
        """)
        rows = cur.fetchall()
        return [row[0] for row in rows]
    finally:
        cur.close()
        pool.putconn(conn)

# this endpoint is for createautomatejob page
@app.get("/busystaff/{job_date}")
def getBusyStaffByDate(job_date: date):
    conn = pool.getconn()
    try:
        cur = conn.cursor()
        cur.execute("""
            SELECT DISTINCT a.empid
            FROM assignments a
            JOIN scheduledjobs sj ON sj.id = a.jobid
            WHERE a.jobdate::date = %s
            AND LOWER(TRIM(sj.status)) IN ('scheduled', 'in progress')
        """, (job_date,))
        rows = cur.fetchall()
        return [row[0] for row in rows]
    finally:
        cur.close()
        pool.putconn(conn)



@app.post("/jobtypes")
def createJT(data: dict):
    conn = pool.getconn()
    name = data.get("jobTypeName")
    minstaff = data.get("minimumStaff")
    skills = data.get("requiredSkills")
    try:
        cur = conn.cursor()
        cur.execute("insert into jobs (type,min_staff) values (%s,%s) RETURNING id",(name,minstaff))
        Id = cur.fetchone()[0]
        clauses = ", ".join(f"{skill} = 1" for skill in skills)
        query = f"UPDATE jobs SET {clauses} WHERE id = %s"
        cur.execute(query,(Id,))
        conn.commit()
        return
    finally:
        cur.close()
        pool.putconn(conn)

@app.delete("/jobtypes/{id}")
def deleteJT(id: str):
    conn = pool.getconn()
    try:
        cur = conn.cursor()
        cur.execute("DELETE FROM jobs WHERE id = %s;", (id,))
        conn.commit()
        return
    finally:
        cur.close()
        pool.putconn(conn)

@app.delete("/jobtypescol/{name}")
def deleteJTcol(name: str):
    conn = pool.getconn()
    try:
        cur = conn.cursor()
        query = sql.SQL("ALTER TABLE jobs DROP COLUMN {}").format(
            sql.Identifier(name)
        )
        cur.execute(query)
        conn.commit()
        return {"message": f"Column '{name}' deleted"}
    finally:
        cur.close()
        pool.putconn(conn)

@app.delete("/employeescol/{name}")
def deleteEMPcol(name: str):
    conn = pool.getconn()
    try:
        cur = conn.cursor()
        query = sql.SQL("ALTER TABLE employees DROP COLUMN {}").format(
            sql.Identifier(name)
        )
        cur.execute(query)
        conn.commit()
        return {"message": f"Column '{name}' deleted"}
    finally:
        cur.close()
        pool.putconn(conn)

@app.put("/jobtypes")
def update_jobtype_skills(data: dict):
    skills = data.get("skills")
    skillI = data.get("skillI")
    name = data.get("name")
    Id = data.get("id")
    min_staff = data.get("min_staff")

    if skills is None or not isinstance(skills, list):
        raise HTTPException(status_code=400, detail="Invalid skills array")

    if skillI is None or not isinstance(skillI, list):
        raise HTTPException(status_code=400, detail="Invalid skill update array")

    if not name:
        raise HTTPException(status_code=400, detail="Provider name is required")

    conn = pool.getconn()
    try:
        cur = conn.cursor()

        set_clauses = []

        for skill in skills:
            set_clauses.append(
                sql.SQL("{} = 0").format(sql.Identifier(skill))
            )

        for skill in skillI:
            set_clauses.append(
                sql.SQL("{} = 1").format(sql.Identifier(skill))
            )

        query = sql.SQL("UPDATE jobs SET {} WHERE id = %s").format(
            sql.SQL(", ").join(set_clauses)
        )

        cur.execute(query, (Id,))
        cur.execute("UPDATE jobs SET type = %s, min_staff = %s WHERE id = %s",(name, min_staff, Id))
        conn.commit()
        return {"message": "Skills updated successfully"}

    finally:
        cur.close()
        pool.putconn(conn)

if __name__ == "__main__":
    import uvicorn
    uvicorn.run(
        "server:app",
        host="0.0.0.0",
        port=int(os.getenv("PORT", 8000)),
        reload=False
    )
